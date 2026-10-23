import openpyxl, json, sys, warnings, datetime, re
from openpyxl.utils import column_index_from_string as ci
warnings.filterwarnings('ignore')
SRC, CHARTS, OUT = sys.argv[1], sys.argv[2], sys.argv[3]
wb = openpyxl.load_workbook(SRC, data_only=True)
charts = {c['chart']: c for c in json.load(open(CHARTS))}

def v(ws, ref):
    return ws[ref].value

def num(x):
    if x is None or x == '': return None
    return float(x)

def r(x, d=4):
    return None if x is None else round(x, d)

D = []  # datasets
def add(**kw):
    kw.setdefault('series', [])
    D.append(kw)
    return kw

SC = 'Statistics Canada and EY analysis'
SP_VIO = 'S&P Global Mobility, Canadian Vehicles in Operation'
SP_NVR = 'S&P Global Mobility, Canadian New Vehicle Registrations'
SP_SBPT = 'S&P Global Mobility, Canadian Sales Based Powertrain forecast'
INRIX = 'AIA Canada and INRIX Travel Trends'
CBS = '2024 AIA Canada Consumer Behaviour Series'
EYMCI = 'EY Mobility Consumer Index'
INTL = 'Statistics Canada, Bank of Canada, Australian Bureau of Statistics, UK ONS, US BEA and EY analysis'

def from_chart(cid, scale=1, nd=4, names=None, axis=None, cat_fn=None):
    c = charts[cid]
    cats = c['series'][0]['cats']
    if cats and isinstance(cats[0], list):  # multilevel: [lvl0, lvl1]
        lv = cats
        if cid == 'chart34.xml':  # quarter + year (year only on first quarter)
            yrs = lv[1]; out = []; y = None
            # lvl1 holds years sparsely in cache; rebuild by counting
            start = 2019
            out = [f"{q} {start + i // 4}" for i, q in enumerate(lv[0])]
            cats = out
        else:
            cats = lv[-1]
    if cat_fn: cats = [cat_fn(x) for x in cats]
    ser = []
    for i, s in enumerate(c['series']):
        vals = [r(x * scale, nd) if x is not None else None for x in s['vals']][:len(cats)]
        e = {'name': (names[i] if names else s['name'].strip()) or 'Value', 'vals': vals}
        if axis and i in axis: e['axis'] = 'right'
        if s['type'] == 'lineChart': e['type'] = 'line'
        ser.append(e)
    return cats, ser

def xl_month(x):
    d = datetime.date(1899, 12, 30) + datetime.timedelta(days=int(float(x)))
    return d.strftime('%b %Y')

# ---------------- SECTOR ----------------
ws = wb['Chapter 1 Data']
cats, ser = from_chart('chart5.xml', scale=0.001, nd=3)
add(id='gdp-trend', topic='sector', ref='Outlook Ch. 1, Fig. 1', title='Auto care sector GDP, Canada, 2012–2023',
    unit='$B', chart='stacked', cats=cats, series=ser, source=SC,
    note='Core = parts manufacturing, wholesale, parts/tire retail and repair. Non-core = gasoline stations and general merchandise retailers.')
cats, ser = from_chart('chart6.xml', nd=2)
add(id='gdp-real-index', topic='sector', ref='Outlook Ch. 1, Fig. 2', title='Real GDP by auto care subsector (2020 = 100), 2020–2023',
    unit='index', chart='line', cats=cats, series=ser, source=SC)

# GDP by province x subsector, 2023 (Table 1)
ws = wb['Sector overview']
provs = [v(ws, f'{c}5') for c in 'CDEFGHIJKLM']
rows = []
for rr in range(6, 13):
    rows.append({'name': v(ws, f'B{rr}').strip(), 'vals': [r(num(v(ws, f'{c}{rr}')) / 1000, 4) for c in 'CDEFGHIJKLM']})
add(id='gdp-province', topic='sector', ref='Sector overview, Table 1', title='Auto care GDP by province and industry, 2023',
    unit='$B', chart='stacked', cats=provs, series=rows, source=SC, orient='h')

# Employment explorers from long data
ws = wb['Employment Data']
emp_prov = {}
for rr in range(5, 148):
    loc, yr, val = v(ws, f'B{rr}'), v(ws, f'C{rr}'), v(ws, f'D{rr}')
    if loc is None: continue
    emp_prov.setdefault(str(loc), {})[str(yr)] = r(num(val), 0)
emp_sub = {}
core = {}
for rr in range(152, 232):
    yr, sub, val, cn = v(ws, f'B{rr}'), v(ws, f'C{rr}'), v(ws, f'D{rr}'), v(ws, f'E{rr}')
    if yr is None or sub is None: continue
    emp_sub.setdefault(sub.strip(), {})[str(yr)] = r(num(val), 0)
    core[sub.strip()] = cn
emp_ps = {}  # prov -> year -> sub -> val
for rr in range(233, ws.max_row + 1):
    loc, yr, sub, val = v(ws, f'B{rr}'), v(ws, f'C{rr}'), v(ws, f'D{rr}'), v(ws, f'E{rr}')
    if loc is None or sub is None or val is None: continue
    emp_ps.setdefault(str(loc), {}).setdefault(str(yr), {})[sub.strip()] = r(num(val), 0)

years = [str(y) for y in range(2013, 2024)]
order = ['ON', 'QC', 'AB', 'BC', 'MB', 'SK', 'NS', 'NB', 'NL', 'PEI', 'YT', 'NWT', 'NU']
add(id='emp-province-trend', topic='sector', ref='Sector overview, Fig. 2', title='Auto care employees by province, 2013–2023',
    unit='people', chart='line', cats=years,
    series=[{'name': p, 'vals': [emp_prov[p].get(y) for y in years]} for p in order if p in emp_prov],
    source=SC, pick='series', pickDefault=['ON', 'QC', 'AB', 'BC'])
subs = list(emp_sub.keys())
add(id='emp-industry-trend', topic='sector', ref='Sector overview, Fig. 1', title='Auto care employees by industry, 2013–2023',
    unit='people', chart='stacked', cats=years,
    series=[{'name': s, 'vals': [emp_sub[s].get(y) for y in years], 'group': core[s]} for s in subs],
    source=SC)
# province x industry for a chosen year (Table 2 explorer)
ps_years = ['2019', '2020', '2021', '2022', '2023']
ps_provs = [p for p in order + ['Territories'] if p in emp_ps]
add(id='emp-prov-industry', topic='sector', ref='Sector overview, Table 2', title='Auto care employees by industry and province',
    unit='people', chart='stacked', orient='h', source=SC, kind='cube',
    cube={'dims': ['year'], 'year': ps_years, 'default': {'year': '2023'},
          'cats': ps_provs, 'series': subs,
          'data': {y: {s: [emp_ps.get(p, {}).get(y, {}).get(s) for p in ps_provs] for s in subs} for y in ps_years}})
cats, ser = from_chart('chart8.xml', nd=0)
add(id='emp-core-trend', topic='sector', ref='Outlook Ch. 1, Fig. 4', title='Employment in core and non-core industries, 2012–2023',
    unit='people', chart='stacked', cats=cats, series=ser, source=SC)
cats, ser = from_chart('chart9.xml', nd=4)
add(id='emp-growth-province', topic='sector', ref='Outlook Ch. 1, Fig. 5', title='Employment growth by province, 2021–2023',
    unit='%', chart='bar', cats=cats, series=ser, source=SC)
cats, ser = from_chart('chart11.xml', nd=0)
add(id='wages', topic='sector', ref='Outlook Ch. 1, Fig. 7', title='Average weekly wages, core and non-core industries, 2016–2023',
    unit='$', chart='bar', cats=cats, series=ser, source=SC)
cats, ser = from_chart('chart12.xml', nd=3, axis=[4])
ser = [s for s in ser if s['name'] != 'Growth rate (right scale)']
for s in ser:
    if s['name'] == 'TOTAL': s['name'] = 'Total'
add(id='technicians', topic='sector', ref='Outlook Ch. 1, Fig. 8', title="Auto care technicians ('000s), Canada, 2019–2024",
    unit="'000", chart='stacked', cats=charts['chart12.xml']['series'][0]['cats'][:len(ser[0]['vals'])], series=ser, source=SC)

# Enterprises
ws = wb['Business Data']
sizes = ['1 to 4 employees', '5 to 9 employees', '10 to 19 employees', '20 to 49 employees', '50 to 99 employees',
         '100 to 199 employees', '200 to 499 employees', '500 plus employees']
short = ['1–4', '5–9', '10–19', '20–49', '50–99', '100–199', '200–499', '500+']
ind = {}
for rr in range(5, 61):
    a, b, c = v(ws, f'B{rr}'), v(ws, f'C{rr}'), v(ws, f'D{rr}')
    if a is None or b not in sizes: continue
    ind.setdefault(a.strip(), {})[b] = num(c)
add(id='enterprises-industry', topic='sector', ref='Sector overview, Table 3', title='Auto care enterprises by size and industry, 2023',
    unit='businesses', chart='stacked', orient='h', cats=list(ind.keys()),
    series=[{'name': short[i], 'vals': [ind[k].get(s) for k in ind]} for i, s in enumerate(sizes)], source=SC)
cm = {}
for rr in range(65, 89):
    a, b, c = v(ws, f'B{rr}'), v(ws, f'C{rr}'), v(ws, f'D{rr}')
    if a is None or b not in sizes: continue
    cm.setdefault(a.strip(), {})[b] = num(c)
add(id='enterprises-collision-mech', topic='sector', ref='Sector overview, Table 4', title='Repair businesses by size: collision, mechanical and other, 2023',
    unit='businesses', chart='stacked', orient='h', cats=list(cm.keys()),
    series=[{'name': short[i], 'vals': [cm[k].get(s) for k in cm]} for i, s in enumerate(sizes)], source=SC)
geo = {}
for rr in range(93, ws.max_row + 1):
    g, s, n, val = v(ws, f'B{rr}'), v(ws, f'C{rr}'), v(ws, f'D{rr}'), v(ws, f'E{rr}')
    if g is None or s not in sizes or n is None: continue
    geo.setdefault(g.strip(), {}).setdefault(re.sub(r'\s+', ' ', n).strip(), {})[s] = num(val)
geos = list(geo.keys())
naics = sorted({n for g in geo.values() for n in g})
cube = {'All Canada': {}}
for g in geos:
    cube[g] = {sz: [geo[g].get(n, {}).get(sz, 0) for n in naics] for sz in sizes}
cube['All Canada'] = {sz: [sum(geo[g].get(n, {}).get(sz, 0) or 0 for g in geos) for n in naics] for sz in sizes}
add(id='enterprises-naics', topic='sector', ref='Sector overview, Table 5', title='Auto care businesses by detailed industry (NAICS) and size',
    unit='businesses', chart='stacked', orient='h', source=SC, kind='cube',
    cube={'dims': ['province'], 'province': ['All Canada'] + geos, 'default': {'province': 'All Canada'},
          'cats': naics, 'series': short,
          'data': {g: {short[i]: cube[g][sz] for i, sz in enumerate(sizes)} for g in cube}})

# Imports (Table 6)
ws = wb['Sector overview']
imp = []; cat = None
for rr in range(123, 203):
    b = v(ws, f'B{rr}'); c = v(ws, f'C{rr}')
    if b is None: continue
    if c is None: cat = b.strip(); continue
    imp.append({'cat': cat, 'name': b.strip(), 'vals': [r(num(v(ws, f'{x}{rr}')), 2) for x in 'CDE']})
groups = {}
for it in imp:
    g = groups.setdefault(it['cat'], [0, 0, 0])
    for i in range(3): g[i] += it['vals'][i]
gsort = sorted(groups.items(), key=lambda kv: -kv[1][2])
add(id='imports', topic='sector', ref='Sector overview, Table 6', title='Auto care product imports by category, 2021–2023',
    unit='$M', chart='bar', orient='h', cats=[k for k, _ in gsort],
    series=[{'name': y, 'vals': [r(g[i], 1) for _, g in gsort]} for i, y in enumerate(['2021', '2022', '2023'])],
    detail={'columns': ['Category', 'Product', '2021', '2022', '2023'], 'rows': [[i['cat'], i['name']] + i['vals'] for i in imp]},
    source=SC)

# ---------------- VEHICLES ----------------
ws = wb['Vehicles in operation']
vio_years = [str(v(ws, f'B{x}')) for x in range(5, 10)]
add(id='vio-national', topic='vehicles', ref='VIO, Tables 1–2', title='Vehicles in operation, passenger cars and light trucks, 2019–2023',
    unit='vehicles', chart='stacked', cats=vio_years,
    series=[{'name': 'Passenger cars', 'vals': [num(v(ws, f'C{x}')) for x in range(5, 10)]},
            {'name': 'Light trucks', 'vals': [num(v(ws, f'D{x}')) for x in range(5, 10)]}],
    source=SP_VIO, note='Forecast: 26.7 million vehicles at year-end 2024 (10.4M passenger cars, 16.3M light trucks) and 31.0 million by year-end 2030 (Table 5).')
regions = [v(ws, f'B{x}') for x in range(23, 28)]
vr = {}
for y, (pc, lt) in {'2021': ('C', 'D'), '2022': ('F', 'G'), '2023': ('I', 'J')}.items():
    vr[y] = {'Passenger cars': [num(v(ws, f'{pc}{x}')) for x in range(23, 28)],
             'Light trucks': [num(v(ws, f'{lt}{x}')) for x in range(23, 28)]}
add(id='vio-region', topic='vehicles', ref='VIO, Table 3', title='Vehicles in operation by region', unit='vehicles',
    chart='stacked', orient='h', source=SP_VIO, kind='cube',
    cube={'dims': ['year'], 'year': ['2021', '2022', '2023'], 'default': {'year': '2023'}, 'cats': regions,
          'series': ['Passenger cars', 'Light trucks'], 'data': vr})
my = [str(v(ws, f'B{x}')) for x in range(34, 78)]
rg = ['BC & Territories', 'Prairies', 'Ontario', 'Quebec', 'Atlantic']
add(id='vio-model-year', topic='vehicles', ref='VIO, Table 4', title='Vehicles in operation by model year and region (Dec. 2023)',
    unit='vehicles', chart='stacked', cats=my,
    series=[{'name': rg[i], 'vals': [num(v(ws, f'{c}{x}')) for x in range(34, 78)]} for i, c in enumerate('CDEFG')],
    source=SP_VIO, note='Model year 2024 is a partial year as of December 2023.')
_mt = [sum(num(v(ws, f'{c}{x}')) for c in 'CDEFG') for x in range(34, 78)]
_old = sum(t for y, t in zip(my, _mt) if int(y) <= 2012); _all = sum(_mt)
_cum = 0
for y, t in zip(reversed(my), reversed(_mt)):
    _cum += t
    if _cum >= _all / 2: _med = y; break
D[-1]['insight'] = f'<strong>{_old/_all*100:.0f}% of vehicles on the road are model year 2012 or older</strong> (11+ years) · half are model year {_med} or older'

ws = wb['New vehicle registrations']
nvr_years = ['2019', '2020', '2021', '2022', '2023']
add(id='nvr-national', topic='vehicles', ref='NVR, Tables 1, 5–6', title='New vehicle registrations, passenger cars and light trucks, 2019–2023',
    unit='vehicles', chart='stacked', cats=nvr_years,
    series=[{'name': 'Passenger cars', 'vals': [num(v(ws, f'C{x}')) for x in range(69, 74)]},
            {'name': 'Light trucks', 'vals': [num(v(ws, f'D{x}')) for x in range(69, 74)]}],
    source=SP_NVR, note='Forecast for year-end 2024: 1.82 million (331,000 passenger cars, 1.49 million light trucks).')

def ranking(rows, colsets, years, clean=lambda s: s):
    out = {}
    for y, cols in zip(years, colsets):
        lst = []
        for x in rows:
            vals = [v(ws, f'{c}{x}') for c in cols]
            if vals[-1] is None: continue
            name = ' '.join(str(a).strip() for a in vals[:-1] if a is not None)
            lst.append([clean(name), num(vals[-1])])
        out[y] = lst
    return out

def tcase(s):
    s = s.strip()
    if s.isupper() and len(s) > 3:
        s = ' '.join(w if w in ('GMC', 'GM', 'CR-V', 'XV', 'EV', 'EUV', 'BMW', 'RAV4', 'CX-5', 'F150', 'SUV') else w.capitalize() for w in s.split())
    s = s.replace('Gm ', 'GM ').replace('Crv', 'CR-V')
    s = re.sub(r'\s+Canada$', '', s)
    s = s.replace('Mach-e', 'Mach-E')
    s = re.sub(r'^KIA\b', 'Kia', s)
    s = s.replace('General Motors (GM)', 'General Motors').replace('GM', 'General Motors') if s in ('GM', 'General Motors (GM)') else s
    return s
add(id='nvr-makes', topic='vehicles', ref='NVR, Table 2', title='New vehicle registrations by make, top 10', unit='vehicles',
    chart='rank', source=SP_NVR, kind='rank', years=nvr_years,
    rank=ranking(range(16, 27), [('B', 'C'), ('D', 'E'), ('F', 'G'), ('H', 'I'), ('J', 'K')], nvr_years, tcase))
add(id='nvr-pc-models', topic='vehicles', ref='NVR, Table 3', title='Top 10 passenger car models by registrations', unit='vehicles',
    chart='rank', source=SP_NVR, kind='rank', years=nvr_years,
    rank=ranking(range(34, 45), [('B', 'C', 'D'), ('E', 'F', 'G'), ('H', 'I', 'J'), ('K', 'L', 'M'), ('R', 'S', 'T')], nvr_years, tcase))
add(id='nvr-lt-models', topic='vehicles', ref='NVR, Table 4', title='Top 10 light truck models by registrations', unit='vehicles',
    chart='rank', source=SP_NVR, kind='rank', years=nvr_years,
    rank=ranking(range(52, 63), [('B', 'C', 'D'), ('E', 'F', 'G'), ('H', 'I', 'J'), ('K', 'L', 'M'), ('R', 'S', 'T')], nvr_years, tcase))
def seg(r1, r2):
    names = [v(ws, f'B{x}').strip() for x in range(r1, r2)]
    return names, [{'name': y, 'vals': [num(v(ws, f'{c}{x}')) for x in range(r1, r2)]} for y, c in zip(nvr_years, 'CDEFG')]
n, s = seg(84, 97)
add(id='nvr-pc-segments', topic='vehicles', ref='NVR, Tables 7, 9, 11', title='Passenger car registrations by segment, 2019–2023',
    unit='vehicles', chart='bar', orient='h', cats=n, series=s, source=SP_NVR, pick='series', pickDefault=['2022', '2023'])
n, s = seg(103, 122)
add(id='nvr-lt-segments', topic='vehicles', ref='NVR, Tables 8, 10, 11', title='Light truck registrations by segment, 2019–2023',
    unit='vehicles', chart='bar', orient='h', cats=n, series=s, source=SP_NVR, pick='series', pickDefault=['2022', '2023'])

# ---------------- EV ----------------
ws = wb['Electric vehicles']
fuels = [('BEV', 38), ('PHEV', 39), ('HEV', 43), ('FCEV', 42), ('ICE', 44)]
add(id='ev-nvr-fuel', topic='ev', ref='EV, Table 3', title='New vehicle registrations by powertrain, 2019–2023',
    unit='vehicles', chart='stacked', cats=nvr_years,
    series=[{'name': f, 'vals': [num(v(ws, f'{c}{x}')) for c in 'CDEFG']} for f, x in fuels], source=SP_NVR,
    pick='series', pickDefault=['BEV', 'PHEV', 'HEV'])
fy = [str(v(ws, f'{c}25')) for c in 'CDEFGHI']
fz = [('BEV', 26), ('PHEV', 27), ('HEV', 31), ('FCEV', 30), ('ICE', 32)]
add(id='ev-forecast', topic='ev', ref='EV, Table 2', title='Sales forecast by powertrain, 2024–2030', unit="'000 vehicles",
    chart='stacked', cats=fy, series=[{'name': f, 'vals': [r(num(v(ws, f'{c}{x}')), 1) for c in 'CDEFGHI']} for f, x in fz],
    source=SP_SBPT)
share_c = ['2019', '2020', '2021', '2022', '2023', '2024', '2025', '2026', '2027', '2028', '2029', '2030']
act = [r(num(v(ws, f'C{x}')), 4) for x in range(72, 77)]
fc = [r(num(v(ws, f'F{x}')), 4) for x in range(72, 79)]
add(id='ev-share', topic='ev', ref='EV, Tables 8–9', title='EVs as a share of new light vehicles: actual and forecast, 2019–2030',
    unit='%', chart='bar', cats=share_c,
    series=[{'name': 'Actual (registrations)', 'vals': act + [None] * 7}, {'name': 'Forecast', 'vals': [None] * 5 + fc}],
    source=SP_NVR + '; ' + SP_SBPT, note='EV = battery electric (BEV) plus plug-in hybrid (PHEV).',
    insight=f'<strong>{act[-1]*100:.1f}% in 2023</strong>, up from {act[0]*100:.1f}% in 2019 · forecast to reach {fc[-1]*100:.1f}% by 2030')
add(id='ev-makes', topic='ev', ref='EV, Table 11', title='EV registrations by make, top 5', unit='vehicles', chart='rank',
    source=SP_NVR, kind='rank', years=nvr_years,
    rank=ranking(range(94, 100), [('B', 'C'), ('D', 'E'), ('F', 'G'), ('H', 'I'), ('J', 'K')], nvr_years, tcase))
add(id='ev-models', topic='ev', ref='EV, Table 12', title='EV registrations by model, top 5', unit='vehicles', chart='rank',
    source=SP_NVR, kind='rank', years=nvr_years,
    rank=ranking(range(105, 111), [('B', 'C', 'D'), ('E', 'F', 'G'), ('H', 'I', 'J'), ('K', 'L', 'M'), ('N', 'O', 'P')], nvr_years, tcase))
launch = []
for top in (4, 8, 12, 17):
    for c in 'BCD':
        brand = v(ws, f'{c}{top}')
        items = [v(ws, f'{c}{x}') for x in range(top + 1, top + 5) if v(ws, f'{c}{x}') and not str(v(ws, f'{c}{x}')).startswith('Source')]
        items = [i for i in items if i not in ('Tesla', 'Ford', 'Volkswagen', 'Hyundai', 'Kia', 'Toyota', 'Jeep', 'Chevrolet', 'Rivian')]
        if brand: launch.append([brand, items])
add(id='ev-launches', topic='ev', ref='EV, Table 1', title='Planned EV launches by brand', kind='cards', cards=launch,
    source='S&P Global Mobility and manufacturer announcements')
cats, ser = from_chart('chart33.xml', nd=3)
add(id='ev-penetration', topic='ev', ref='Outlook Ch. 4, Fig. 27', title='EV market penetration (consumer survey), 2021–2023',
    unit='%', chart='stacked', cats=cats, series=ser, source=EYMCI,
    insight=f'<strong>EV penetration reached {ser[0]["vals"][-1]*100:.0f}% in {cats[-1]}</strong>, up from {ser[0]["vals"][0]*100:.0f}% in {cats[0]}')
cats, ser = from_chart('chart36.xml', nd=3)
add(id='ev-intent', topic='ev', ref='Outlook Ch. 4, Fig. 28', title='Purchase intent: EV vs. other vehicles, 2021–2023',
    unit='%', chart='bar', cats=cats, series=ser, source=EYMCI)
cats, ser = from_chart('chart37.xml', nd=3, names=['Share of respondents'])
add(id='ev-motivators', topic='ev', ref='Outlook Ch. 4, Fig. 29', title='Motivators for buying an EV', unit='%', chart='bar',
    orient='h', cats=cats[::-1], series=[{'name': ser[0]['name'], 'vals': ser[0]['vals'][::-1]}], source=EYMCI,
    note='The survey label reads "Total cost of ownership is lower than EVs" as published; it likely means lower than gas vehicles.')
cats, ser = from_chart('chart38.xml', nd=3, names=['Share of respondents'])
add(id='ev-deterrents', topic='ev', ref='Outlook Ch. 4, Fig. 30', title='Key deterrents for potential EV buyers', unit='%', chart='bar',
    orient='h', cats=cats[::-1], series=[{'name': ser[0]['name'], 'vals': ser[0]['vals'][::-1]}], source=EYMCI)

# ---------------- DEMAND ----------------
ws = wb['Demand analysis']
rows = list(range(33, 41))
names = [v(ws, f'B{x}').strip() for x in rows]
add(id='market-size', topic='demand', ref='Demand analysis, Table 1', title='Size of the core auto care market by subsector, 2020–2022',
    unit='$B', chart='bar', orient='h', cats=names,
    series=[{'name': y, 'vals': [r(num(v(ws, f'{c}{x}')), 3) for x in rows]} for y, c in zip(['2020', '2021', '2022'], 'CDE')],
    source=SC, pick='series', pickDefault=['2020', '2022'], note='Collision, mechanical and other repair are a breakdown of automotive repair and maintenance. The 2021 column in the source is much lower than 2020 and 2022 for every line (total $10.7B vs. $30.7B and $35.5B); treat 2021 with caution until confirmed with EY.')
rows = list(range(48, 53))
add(id='market-region', topic='demand', ref='Demand analysis, Table 2 / Fig. 2', title='Size of the core auto care market by region, 2020–2022',
    unit='$B', chart='bar', orient='h', cats=[v(ws, f'B{x}').strip() for x in rows],
    series=[{'name': y, 'vals': [r(num(v(ws, f'{c}{x}')), 3) for x in rows]} for y, c in zip(['2020', '2021', '2022'], 'CDE')],
    source='Statistics Canada, AIA Canada and EY analysis', pick='series', pickDefault=['2020', '2022'], note='Regional split relies on AIA Canada estimates. Same 2021 caution as the subsector table.')
rep = []
for x in range(76, 261):
    rep.append([v(ws, f'B{x}').strip()] + [r(num(v(ws, f'{c}{x}')), 4) for c in 'CDEFG'])
add(id='replacement-rates', topic='demand', ref='Demand analysis, Table 3', title='Average annual replacement rate by product and vehicle age, 2023',
    kind='table', unit='per vehicle per year', columns=['Product', 'New–5 yrs', '6–10 yrs', '11–15 yrs', '16+ yrs', 'All vehicles'],
    rows=rep, source='IMR Inc. All rights reserved.', note='Read 0.25 as 25 replacements per 100 vehicles per year (oil changes exceed 1.0 because vehicles get more than one a year).')
cats, ser = from_chart('chart20.xml', nd=3)
add(id='collision-mech-revenue', topic='demand', ref='Outlook Ch. 2, Fig. 17', title='Collision and mechanical sector revenues, 2020–2022',
    unit='$B', chart='bar', cats=cats, series=ser, source=SC)
cats, ser = from_chart('chart13.xml', nd=3)
add(id='vkt-index', topic='demand', ref='Outlook Ch. 2, Fig. 9', title='Total vehicle kilometres travelled index, Canada, by quarter',
    unit='index', chart='bar', cats=cats, series=ser, source=INRIX)
cats, ser = from_chart('chart21.xml', nd=3)
add(id='trip-duration', topic='demand', ref='Outlook Ch. 2, Fig. 12', title='Average trip duration index, Canada, by quarter',
    unit='index', chart='line', cats=cats, series=ser, source=INRIX)
cats, ser = from_chart('chart14.xml', nd=1, names=['Vehicle kilometres travelled'])
add(id='vkt-annual', topic='demand', ref='Outlook Ch. 2, Fig. 10', title='Change in total vehicle kilometres travelled (2020 = 100)',
    unit='index', chart='bar', cats=cats, series=ser, source=INRIX)
cats, ser = from_chart('chart15.xml', nd=3, names=['Travel distance index'])
add(id='vkt-province', topic='demand', ref='Outlook Ch. 2, Fig. 11', title='Travel distance by province, 2023', unit='index',
    chart='bar', cats=cats, series=ser, source=INRIX)
cats, ser = from_chart('chart16.xml', nd=3, names=['Share of consumers'])
add(id='repair-channel', topic='demand', ref='Outlook Ch. 2, Fig. 13', title='Where Canadians get maintenance and repairs done, 2023',
    unit='%', chart='bar', orient='h', cats=cats, series=ser, source=CBS)
cats, ser = from_chart('chart17.xml', nd=3)
cats = ['Support the Right to Repair requirement to equip all vehicles with telematics', 'Consumers should be able to get vehicles serviced where they choose',
        'Manufacturers should be required by law to share repair data', 'Less likely to buy a brand that restricts repair access']
add(id='r2r-opinions', topic='demand', ref='Outlook Ch. 2, Fig. 14', title='Consumer views on the Right to Repair, 2023', unit='%',
    chart='stacked', orient='h', cats=cats, series=ser, source=CBS, note='Statements are shortened; see the Outlook Study for the full survey wording.')
cats, ser = from_chart('chart19.xml', nd=2, axis=[1])
add(id='collisions', topic='demand', ref='Outlook Ch. 2, Fig. 16', title='Vehicles involved in collisions and fatality rate, 2009–2021',
    unit='vehicles', chart='combo', cats=cats, series=ser, source='Statistics Canada, Transport Canada and EY analysis',
    rightUnit='per 10,000 vehicles')
cats, ser = from_chart('chart18.xml', nd=3, axis=[1])
add(id='repair-revenue', topic='demand', ref='Outlook Ch. 2, Fig. 15', title='Repair and maintenance: operating revenue index (2017 = 100) and real GDP growth',
    unit='index', chart='combo', cats=cats, series=ser, source='ISED Canada and EY analysis', rightUnit='%')

# ---------------- ECONOMY ----------------
ws = wb['National Accounts Data']
gp = {}
for rr in range(62, ws.max_row + 1):
    y, p, val = v(ws, f'B{rr}'), v(ws, f'C{rr}'), v(ws, f'D{rr}')
    if y is None or p is None or val is None: continue
    try: int(y)
    except: continue
    gp.setdefault(re.sub(r'\s+\d+$','',p.strip()), {})[str(y)] = r(num(val), 4)
gy = sorted({y for d in gp.values() for y in d})
add(id='gdp-growth-province', topic='economy', ref='Demand analysis, Fig. 1', title='Real GDP growth by province (year over year)',
    unit='%', chart='bar', source='Statistics Canada Table 36-10-0222-01 and EY analysis', kind='cube', orient='h',
    cube={'dims': ['year'], 'year': gy, 'default': {'year': gy[-1]}, 'cats': list(gp.keys()), 'series': ['Growth'],
          'data': {y: {'Growth': [gp[p].get(y) for p in gp]} for y in gy}})
cats, ser = from_chart('chart39.xml', nd=4, cat_fn=xl_month)
add(id='cpi', topic='economy', ref='Outlook Ch. 4, Fig. 31', title='Consumer price index, year-over-year change', unit='%',
    chart='line', cats=cats, series=ser, source='Statistics Canada, Bank of Canada and EY analysis')
cats, ser = from_chart('chart35.xml', nd=4, cat_fn=xl_month)
add(id='ippi', topic='economy', ref='Outlook Ch. 4, Fig. 33', title='Industrial product price index, year-over-year change', unit='%',
    chart='line', cats=cats, series=ser, source='Statistics Canada, Bank of Canada and EY analysis', pick='series',
    pickDefault=['Total, Industrial product price index', 'Motor vehicle engines and motor vehicle parts', 'Motor vehicle brakes and brake systems'])
cats, ser = from_chart('chart34.xml', nd=4)
add(id='inflation-expectations', topic='economy', ref='Outlook Ch. 4, Fig. 32', title='Consumer inflation expectations by time horizon',
    unit='%', chart='line', cats=cats, series=ser, source='Bank of Canada and EY analysis')
for cid, idn, ttl, rf in [('chart40.xml', 'ma-domestic', 'Mergers and acquisitions: Canadian buyers of Canadian companies', 'Fig. 34'),
                          ('chart41.xml', 'ma-outbound', 'Mergers and acquisitions: Canadian buyers of foreign companies', 'Fig. 35'),
                          ('chart42.xml', 'ma-inbound', 'Mergers and acquisitions: foreign buyers of Canadian companies', 'Fig. 36')]:
    cats, ser = from_chart(cid, nd=0)
    add(id=idn, topic='economy', ref='Outlook Ch. 4, ' + rf, title=ttl + ', by quarter', unit='deals', chart='stacked',
        cats=cats, series=ser, source='Mergermarket and EY analysis',
        insight=f'<strong>{int(sum(sum(x or 0 for x in s_["vals"]) for s_ in ser))} deals</strong> from {cats[0]} to {cats[-1]} · {int(sum(x or 0 for x in ser[0]["vals"]))} in core auto care industries')

# ---------------- INTERNATIONAL ----------------
for cid, idn, ttl, rf, unit, ch in [
    ('chart23.xml', 'intl-gdp', 'Real GDP by country (2020 = 100)', 'Fig. 18', 'index', 'line'),
    ('chart24.xml', 'intl-sector-gdp', 'Auto care sector GDP as a share of total GDP, 2023', 'Fig. 19', '%', 'bar'),
    ('chart25.xml', 'intl-household', 'Vehicle operating costs as a share of household spending, 2022', 'Fig. 20', '%', 'bar'),
    ('chart26.xml', 'intl-workforce', 'Auto care workforce as a share of total employment, 2023', 'Fig. 21', '%', 'bar'),
    ('chart27.xml', 'intl-registrations', 'Vehicle registrations per 1,000 people, 2023', 'Fig. 22', 'per 1,000', 'bar'),
    ('chart28.xml', 'intl-vkt', 'Total vehicle kilometres travelled (2020 = 100)', 'Fig. 23', 'index', 'line'),
    ('chart29.xml', 'intl-core', 'Core vs. non-core share of auto care employment, 2023', 'Fig. 24', '%', 'stacked'),
    ('chart30.xml', 'intl-sales', 'New vehicle sales per 1,000 people, 2019–2023', 'Fig. 25', 'per 1,000', 'bar'),
    ('chart31.xml', 'intl-ev', 'EV share of new vehicle sales, 2019–2023', 'Fig. 26', '%', 'bar')]:
    cats, ser = from_chart(cid, nd=4)
    for s in ser:
        if not s['name'] or s['name'] == 'Value': s['name'] = 'Value'
    add(id=idn, topic='intl', ref='Outlook Ch. 3, ' + rf, title=ttl, unit=unit, chart=ch, cats=cats, series=ser,
        source=INTL, highlight='Canada')
    if idn == 'intl-core':
        D[-1]['insight'] = '<strong>Canada: ' + f'{ser[0]["vals"][0]*100:.1f}% core, {ser[1]["vals"][0]*100:.1f}% non-core</strong> · the UK has the highest core share at {ser[0]["vals"][2]*100:.1f}%'

# ---------------- DEFINITIONS ----------------
ws = wb['A.1.3 Electric-hybrid vehicles']
evdef = [[v(ws, f'B{x}'), v(ws, f'C{x}'), v(ws, f'D{x}')] for x in (6, 7, 9, 10)]
ws = wb['A.1.1 Aftermarket sector']
sector = []
cur = None
for x in range(5, ws.max_row + 1):
    a, b, c = v(ws, f'B{x}'), v(ws, f'C{x}'), v(ws, f'D{x}')
    if a is not None and b is not None:
        cur = [str(a).strip(), b.strip(), []]; sector.append(cur)
    if c and cur and not str(a or '').startswith('Source'):
        cur[2].append(str(c).strip())
for s in sector:
    s[2] = s[2][:6]
add(id='definitions', topic='defs', title='Definitions', kind='defs', ev=evdef, sector=[s for s in sector if len(s[0]) >= 4][:60],
    source='Library of Parliament; Statistics Canada NAICS 2022; AIA Canada')

KPI = {
    'gdp': 43874.6, 'emp': 507120, 'ent': 49268, 'vio': 25104996, 'nvr': 1599246, 'evshare': 0.1166,
}
json.dump({'datasets': D, 'kpi': KPI, 'built': '2025-01-08'}, open(OUT, 'w'), separators=(',', ':'), ensure_ascii=False)
print(len(D), 'datasets')
for d in D:
    print(d['id'], d.get('kind', d.get('chart')), len(d.get('cats', [])) if d.get('cats') else '', [s['name'] for s in d['series']][:6] if d['series'] else '')
